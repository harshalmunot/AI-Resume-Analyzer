# extracts text from resume files
# supports PDF (using pdfplumber) and DOCX (using python-docx)
# note: image-only/scanned PDFs won't work (would need OCR)

import io
import pdfplumber
from docx import Document


def extract_text(stream, filename):
    """
    read text from an uploaded resume file
    """
    data = stream.read()
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        text = read_pdf(data)
    elif ext == "docx":
        text = read_docx(data)
    else:
        raise ValueError(f"unsupported file type: {ext}")

    # clean up whitespace before returning
    return clean_text(text)


def read_pdf(data):
    # go through each page and grab text
    all_text = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            all_text.append(page_text)
    return "\n".join(all_text)


def read_docx(data):
    doc = Document(io.BytesIO(data))
    parts = []

    # regular paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    # also grab text from tables (some resumes use tables for layout)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def clean_text(text):
    # strip empty lines and extra whitespace
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)
